"""Regression tests for evidence source retention, crop pixels, and placement."""

import copy
import io
import unittest
import zipfile

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Pt

from ppt_builder import AdvancedPPTMaker, PPT_LAYOUT_VERSION


def DetailedPPTMaker(data, assets=None, analyses=None):
    """Retain the legacy layout contract when explicitly requested."""
    return AdvancedPPTMaker({**data, "presentation_mode": "detailed"}, assets, analyses)


def asset(asset_id, color="red", width=800, height=400):
    image = Image.new("RGB", (width, height), color)
    # Different halves make the crop itself independently verifiable.
    image.paste("blue", (width // 2, 0, width, height))
    output = io.BytesIO()
    image.save(output, format="PNG")
    return {"id": asset_id, "name": "patient_name_123_report.png", "image_bytes": output.getvalue(),
            "width": width, "height": height}


def analysis(asset_id, category="pathology", regions=None, findings=None):
    return {"asset_id": asset_id, "category": category, "date": "2026-10-01",
            "title": "检测结果", "findings": findings or ["样本检测结果。"], "transcription": "",
            "regions": regions or [], "warnings": []}


def link(asset_id, section, treatment_index=None, region_ids=None):
    return {"asset_id": asset_id, "region_ids": region_ids or [], "section": section,
            "treatment_index": treatment_index, "caption": "报告图示"}


def slide_text(slide):
    return "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)


class EvidenceDeckTests(unittest.TestCase):
    def test_real_cropped_pixels_and_aspect_ratio_are_embedded(self):
        source = asset("IMG001")
        region = {"id": "r1", "label": "病理切片", "box": [500, 0, 1000, 1000]}
        maker = DetailedPPTMaker({"evidence_links": [link("IMG001", "baseline")]},
                                 [source], [analysis("IMG001", regions=[region])])
        stream = maker.build()
        with zipfile.ZipFile(stream) as archive:
            media = [name for name in archive.namelist() if name.startswith("ppt/media/")]
            self.assertEqual(len(media), 1)
            with Image.open(io.BytesIO(archive.read(media[0]))) as embedded:
                self.assertEqual(embedded.size, (400, 400))
                self.assertEqual(embedded.convert("RGB").getpixel((200, 200)), (0, 0, 255))
        stream.seek(0)
        deck = Presentation(stream)
        picture = next(shape for shape in deck.slides[2].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)
        self.assertAlmostEqual(picture.width / picture.height, 1.0, places=5)
        self.assertEqual(picture.crop_left, 0)
        self.assertEqual(picture.crop_right, 0)
        self.assertIn("资料 01", slide_text(deck.slides[2]))
        self.assertIn("2026-10-01", slide_text(deck.slides[2]))
        self.assertIn("报告图示", slide_text(deck.slides[2]))
        self.assertNotIn("patient_name", "\n".join(slide_text(slide) for slide in deck.slides))

    def test_sections_and_treatment_indices_control_insert_order(self):
        identifiers = ["base", "tx0", "tx1", "current", "fallback", "unmapped"]
        data = {"treatments": [{"phase": "辅助治疗", "regimen": "方案A"}, {"phase": "第二阶段"}],
                "current_admission": {"exams": ["检查"]},
                "evidence_links": [link("base", "baseline"), link("tx0", "treatment", 0),
                                   link("tx1", "treatment", 1), link("current", "current_admission"),
                                   link("fallback", "treatment", 200)]}
        maker = DetailedPPTMaker(data, [asset(item) for item in identifiers],
                                 [analysis(item) for item in identifiers])
        deck = Presentation(maker.build())
        texts = [slide_text(slide) for slide in deck.slides]
        index = {}
        for entry in maker.evidence_manifest:
            index.setdefault(entry["asset_id"], entry["slide_number"] - 1)
        self.assertIn("基线资料", texts[index["base"] - 1])
        self.assertIn("治疗经过：辅助治疗", texts[index["tx0"] - 1])
        self.assertIn("治疗经过：第二阶段", texts[index["tx1"] - 1])
        self.assertIn("本次入院", texts[index["current"] - 1])
        self.assertLess(index["current"], index["fallback"])
        self.assertLess(index["fallback"], index["unmapped"])
        self.assertIn("病例思考与总结", texts[-1])
        self.assertEqual(maker.evidence_manifest[-2]["section"], "appendix")

    def test_unknown_regions_fall_back_to_full_image_and_missing_analysis_is_retained(self):
        sources = [asset("invalid"), asset("no_analysis")]
        regions = [{"id": "actual", "label": "结果", "box": [0, 0, 500, 1000]}]
        maker = DetailedPPTMaker({"evidence_links": [link("invalid", "baseline", region_ids=["invented"]),
                                                    link("missing_source", "baseline")]},
                                 sources, [analysis("invalid", regions=regions)])
        deck = Presentation(maker.build())
        self.assertEqual({entry["asset_id"] for entry in maker.evidence_manifest}, {"invalid", "no_analysis"})
        self.assertTrue(all(entry["section"] == "appendix" for entry in maker.evidence_manifest))
        pictures = [shape for slide in deck.slides for shape in slide.shapes
                    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
        self.assertEqual(len(pictures), 2)
        for picture in pictures:
            with Image.open(io.BytesIO(picture.image.blob)) as image:
                self.assertEqual(image.size, (800, 400))
        self.assertTrue(maker.warnings)

    def test_exclusions_and_absent_current_admission_are_respected(self):
        maker = DetailedPPTMaker({"excluded_asset_ids": ["skip"],
                                 "evidence_links": [link("skip", "baseline"),
                                                    link("keep", "current_admission")]},
                                 [asset("skip"), asset("keep")], [])
        maker.build()
        self.assertEqual([entry["asset_id"] for entry in maker.evidence_manifest], ["keep"])
        self.assertEqual(maker.evidence_manifest[0]["section"], "appendix")

    def test_no_mutation_or_implicit_medical_reclassification(self):
        data = {"treatments": [{"phase": "辅助治疗", "regimen": "X"}],
                "timeline_events": [{"date": "2026", "phase": "辅助", "event": "记录"}]}
        original = copy.deepcopy(data)
        sources = [asset("source")]
        analyses = [analysis("source")]
        original_analyses = copy.deepcopy(analyses)
        maker = DetailedPPTMaker(data, sources, analyses)
        deck = Presentation(maker.build())
        self.assertEqual(data, original)
        self.assertEqual(analyses, original_analyses)
        self.assertIn("治疗经过：辅助治疗", "\n".join(slide_text(slide) for slide in deck.slides))

    def test_more_than_two_regions_and_long_findings_are_paginated_without_loss(self):
        regions = [{"id": f"r{i}", "label": f"区域{i}", "box": [i * 200, 0, (i + 1) * 200, 1000]}
                   for i in range(5)]
        finding = "病理所见" * 350 + "末尾标记END"
        maker = DetailedPPTMaker({}, [asset("long")], [analysis("long", regions=regions, findings=[finding])])
        deck = Presentation(maker.build())
        for slide in deck.slides:
            self.assertLessEqual(sum(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in slide.shapes), 2)
        self.assertEqual(sum(len(entry["region_ids"]) for entry in maker.evidence_manifest if entry["kind"] == "images"), 5)
        self.assertGreater(sum(entry["kind"] == "findings" for entry in maker.evidence_manifest), 1)
        text = "\n".join(slide_text(slide) for slide in deck.slides)
        self.assertIn("末尾标记END", text.replace("\n", "").replace("\v", ""))
        self.assertEqual(text.count("病理所见"), 350)
        for entry in maker.evidence_manifest:
            slide = deck.slides[entry["slide_number"] - 1]
            for shape in slide.shapes:
                self.assertGreaterEqual(shape.left, 0)
                self.assertGreaterEqual(shape.top, 0)
                self.assertLessEqual(shape.left + shape.width, deck.slide_width)
                self.assertLessEqual(shape.top + shape.height, deck.slide_height)

    def test_timeline_retains_all_events_and_cards_fit_the_slide(self):
        events = [{"date": f"2026-{i + 1:02d}", "event": f"事件{i + 1:02d}"} for i in range(13)]
        deck = Presentation(DetailedPPTMaker({"timeline_events": events}).build())
        timelines = [slide for slide in deck.slides if "全病程时间轴概览" in slide_text(slide)]
        self.assertEqual(len(timelines), 3)
        texts = "\n".join(slide_text(slide) for slide in timelines)
        self.assertTrue(all(event["event"] in texts for event in events))
        for slide in timelines:
            for shape in slide.shapes:
                self.assertGreaterEqual(shape.left, 0)
                self.assertGreaterEqual(shape.top, 0)
                self.assertLessEqual(shape.left + shape.width, deck.slide_width)
                self.assertLessEqual(shape.top + shape.height, deck.slide_height)

    def test_existing_sections_paginate_long_content_without_loss(self):
        body = "重复检查结果。" * 220
        phase = "长期治疗阶段" * 25 + "PHASE_END"
        values = {key: body + key + "_END" for key in ("BASE", "REGIMEN", "IMAGING", "MARKERS", "EXAMS", "PLAN", "SUMMARY", "DISCUSSION")}
        data = {"baseline": {"key_exams": values["BASE"]},
                "treatments": [{"phase": phase, "regimen": values["REGIMEN"], "imaging": values["IMAGING"], "markers": values["MARKERS"]}],
                "current_admission": {"exams": [values["EXAMS"]], "plan": [values["PLAN"]]},
                "summary": {"highlights": [values["SUMMARY"]], "discussion": [values["DISCUSSION"]]},
                "evidence_links": [link("treatment_image", "treatment", 0)]}
        original = copy.deepcopy(data)
        maker = DetailedPPTMaker(data, [asset("treatment_image")], [analysis("treatment_image")])
        deck = Presentation(maker.build())
        texts = [slide_text(slide).replace("\n", "").replace("\v", "") for slide in deck.slides]
        combined = "".join(shape.text.replace("\n", "").replace("\v", "")
                           for slide in deck.slides for shape in slide.shapes
                           if shape.has_text_frame and shape.top >= 914400)
        for value in values.values():
            self.assertIn(value, combined)
        self.assertIn(phase, combined)
        image_index = maker.evidence_manifest[0]["slide_number"] - 1
        self.assertIn("MARKERS_END", texts[image_index - 1])
        last_evidence_index = max(entry["slide_number"] for entry in maker.evidence_manifest) - 1
        self.assertIn("本次入院", texts[last_evidence_index + 1])
        self.assertEqual(data, original)
        self.assertGreater(len(deck.slides), 15)
        for slide in deck.slides:
            for shape in slide.shapes:
                self.assertGreaterEqual(shape.left, 0)
                self.assertGreaterEqual(shape.top, 0)
                self.assertLessEqual(shape.left + shape.width, deck.slide_width)
                self.assertLessEqual(shape.top + shape.height, deck.slide_height)

    def test_unhashable_source_references_fall_back_without_crashing(self):
        sources = [asset("valid"), {"id": ["invalid"], "image_bytes": b"invalid"}]
        maker = DetailedPPTMaker({"excluded_asset_ids": [["invalid"]],
                                 "evidence_links": [{"asset_id": ["bad"]},
                                                    {"asset_id": "valid", "section": {"bad": True}, "region_ids": [["bad"]]}]},
                                 sources, [{"asset_id": ["bad"]}, analysis("valid")])
        maker.build()
        self.assertEqual(len(maker.evidence_manifest), 1)
        self.assertEqual(maker.evidence_manifest[0]["asset_id"], "valid")
        self.assertEqual(maker.evidence_manifest[0]["section"], "appendix")
        self.assertTrue(maker.warnings)

    def test_original_upload_source_index_is_preserved(self):
        retained = asset("retained")
        retained["source_index"] = 7
        maker = DetailedPPTMaker({}, [retained], [])
        deck = Presentation(maker.build())
        self.assertEqual(maker.evidence_manifest[0]["source_label"], "资料 07")
        self.assertIn("资料 07", "\n".join(slide_text(slide) for slide in deck.slides))

    def test_image_caption_moves_a_whole_sentence_instead_of_cutting_it(self):
        finding = "连续完整检查结果" * 14 + "，报告保留完整结论。"
        item = analysis("report", findings=[finding, "第二个完整结论。"])
        maker = DetailedPPTMaker({"evidence_links": [link("report", "baseline")]}, [asset("report")], [item])
        deck = Presentation(maker.build())
        image_page = next(entry for entry in maker.evidence_manifest if entry["kind"] == "images")
        image_text = slide_text(deck.slides[image_page["slide_number"] - 1])
        self.assertNotIn("连续完整检查结果", image_text)
        result_text = "".join(shape.text.replace("\n", "").replace("\v", "")
                              for entry in maker.evidence_manifest if entry["kind"] == "findings"
                              for shape in deck.slides[entry["slide_number"] - 1].shapes
                              if shape.has_text_frame and 914400 <= shape.top < 6300000)
        self.assertIn(finding, result_text)
        self.assertIn("第二个完整结论。", result_text)

    def test_portrait_report_uses_available_height_without_stretching(self):
        maker = DetailedPPTMaker({}, [asset("portrait", width=500, height=1500)], [analysis("portrait")])
        deck = Presentation(maker.build())
        picture = next(shape for slide in deck.slides for shape in slide.shapes
                       if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)
        self.assertGreaterEqual(picture.height, 5 * 914400)
        self.assertAlmostEqual(picture.width / picture.height, 1 / 3, places=5)
        self.assertEqual(picture.crop_top + picture.crop_bottom + picture.crop_left + picture.crop_right, 0)

    def test_wide_text_regions_stack_at_full_slide_width(self):
        regions = [{"id": "r1", "label": "诊断", "box": [0, 0, 1000, 100]},
                   {"id": "r2", "label": "免疫组化", "box": [0, 100, 1000, 200]}]
        maker = DetailedPPTMaker({}, [asset("strips")], [analysis("strips", regions=regions)])
        deck = Presentation(maker.build())
        pictures = [shape for slide in deck.slides for shape in slide.shapes
                    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
        self.assertEqual(len(pictures), 2)
        self.assertTrue(all(picture.width >= 11 * 914400 for picture in pictures))
        self.assertLess(pictures[0].top + pictures[0].height, pictures[1].top)
        self.assertTrue(all(picture.crop_top + picture.crop_bottom + picture.crop_left + picture.crop_right == 0
                            for picture in pictures))

    def test_explicit_report_group_keeps_both_images_and_source_references(self):
        sources = [asset("report_a", "red", 900, 600), asset("report_b", "green", 900, 600)]
        data = {"evidence_links": [link("report_a", "baseline"), link("report_b", "baseline")],
                "evidence_groups": [{"group_id": "report-123", "asset_ids": ["report_a", "report_b"],
                                     "reason": "matching_report_content", "section": "baseline", "treatment_index": None}]}
        maker = DetailedPPTMaker(data, sources, [analysis("report_a"), analysis("report_b")])
        deck = Presentation(maker.build())
        images = [entry for entry in maker.evidence_manifest if entry["kind"] == "images"]
        self.assertEqual({entry["asset_id"] for entry in images}, {"report_a", "report_b"})
        self.assertEqual(len({entry["slide_number"] for entry in images}), 1)
        slide = deck.slides[images[0]["slide_number"] - 1]
        self.assertEqual(sum(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in slide.shapes), 2)
        self.assertIn("资料 01", slide_text(slide))
        self.assertIn("资料 02", slide_text(slide))
        del data["evidence_groups"]
        ungrouped = DetailedPPTMaker(data, sources, [analysis("report_a"), analysis("report_b")])
        ungrouped.build()
        self.assertEqual(len({entry["slide_number"] for entry in ungrouped.evidence_manifest
                              if entry["kind"] == "images"}), 2)

    def test_genetic_result_tables_use_one_region_per_slide(self):
        regions = [{"id": "table_1", "label": "I/II 类变异", "box": [0, 0, 1000, 500]},
                   {"id": "table_2", "label": "III 类变异", "box": [0, 500, 1000, 1000]}]
        maker = DetailedPPTMaker({}, [asset("genetic_tables", width=1200, height=1200)],
                                 [analysis("genetic_tables", "genetics", regions=regions)])
        deck = Presentation(maker.build())
        images = [entry for entry in maker.evidence_manifest if entry["kind"] == "images"]
        self.assertEqual(len(images), 2)
        self.assertTrue(all(len(entry["region_ids"]) == 1 for entry in images))
        for entry in images:
            pictures = [shape for shape in deck.slides[entry["slide_number"] - 1].shapes
                        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
            self.assertEqual(len(pictures), 1)
            self.assertGreater(pictures[0].width, 8 * 914400)
            self.assertEqual(pictures[0].crop_top + pictures[0].crop_bottom +
                             pictures[0].crop_left + pictures[0].crop_right, 0)
        self.assertEqual(PPT_LAYOUT_VERSION, "2026-10-22.5")

    def test_internal_phase_enums_are_rendered_as_chinese(self):
        events = [{"date": "2024", "phase": "baseline", "event": "诊断检查"},
                  {"date": "2025", "phase": "treatment", "event": "治疗记录"},
                  {"date": "2026", "phase": "current_admission", "event": "本次计划"}]
        deck = Presentation(DetailedPPTMaker({"timeline_events": events,
                                              "treatments": [{"phase": "treatment"}]}).build())
        text = "\n".join(slide_text(slide) for slide in deck.slides)
        self.assertIn("【基线评估】", text)
        self.assertIn("【治疗阶段】", text)
        self.assertIn("【本次评估】", text)
        self.assertNotIn("baseline", text)
        self.assertNotIn("current_admission", text)
        self.assertNotIn("treatment", text)

    def test_manual_wrapping_keeps_assay_identifiers_and_short_drug_names_intact(self):
        for token in ("CK5/6(-)", "PD-L1(E1L3N)(-)", "卡度尼利单抗", "卡度尼利", "替莫唑胺"):
            text = "检查结果记录" + token + "，继续记录。"
            lines = AdvancedPPTMaker._wrap_lines(text, max_units=20)
            self.assertEqual("".join(lines), text)
            self.assertTrue(any(token in line for line in lines), token)


class CompactDeckTests(unittest.TestCase):
    def test_chinese_closing_punctuation_never_becomes_an_orphan_line(self):
        for text, limit in [("期间联合两次TAE。", 15), ("术后补液、护肝。", 8),
                            ("CK5/6(-)，继续记录。", 8)]:
            lines = AdvancedPPTMaker._wrap_lines(text, limit)
            self.assertEqual("".join(lines), text)
            self.assertFalse(any(line.startswith(("。", "，", "、")) for line in lines[1:]), lines)
        self.assertTrue(any("TAE。" in line for line in AdvancedPPTMaker._wrap_lines("期间联合两次TAE。", 15)))

    def test_explicit_plan_groups_three_sources_and_keeps_original_crop_pixels(self):
        sources = [asset(f"a{i}", width=800, height=400) for i in range(3)]
        records = [analysis(f"a{i}", regions=[{"id": "right", "label": "病理", "box": [500, 0, 1000, 1000]}])
                   for i in range(3)]
        spec = {"type": "content", "title": "病理检查汇总", "section": "baseline",
                "sections": [{"heading": "诊断", "text": "三个标本的结果分别列示。"}],
                "images": [{"asset_id": f"a{i}", "region_id": "right", "label": f"标本{i + 1}"} for i in range(3)]}
        maker = AdvancedPPTMaker({"presentation_plan": {"slides": [spec]}}, sources, records)
        deck = Presentation(maker.build())
        self.assertEqual(len(deck.slides), 1)
        pictures = [shape for shape in deck.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
        self.assertEqual(len(pictures), 3)
        for picture in pictures:
            with Image.open(io.BytesIO(picture.image.blob)) as image:
                self.assertEqual(image.size, (400, 400))
                self.assertEqual(image.getpixel((100, 100)), (0, 0, 255))
        self.assertEqual({entry["asset_id"] for entry in maker.evidence_manifest}, {"a0", "a1", "a2"})
        self.assertTrue(all(entry["region_ids"] == ["right"] for entry in maker.evidence_manifest))
        self.assertTrue(all(entry["section"] == "baseline" for entry in maker.evidence_manifest))

    def test_editable_line_table_preserves_explicit_labels_without_inferring_numbers(self):
        data = {"treatments": [{"line": 1, "duration": "2020", "regimen": "方案甲", "response": "部分缓解"},
                                {"line_label": "二线", "duration": "2021", "regimen": "方案乙", "imaging": "稳定"},
                                {"phase": "维持治疗", "duration": "2022", "regimen": "方案丙"}]}
        deck = Presentation(AdvancedPPTMaker(data).build())
        tables = [shape.table for slide in deck.slides for shape in slide.shapes if shape.has_table]
        self.assertEqual(len(tables), 1)
        cells = [cell.text for row in tables[0].rows for cell in row.cells]
        self.assertIn("第1线", cells)
        self.assertIn("二线", cells)
        self.assertIn("维持治疗", cells)
        self.assertNotIn("三线", cells)
        self.assertIn("部分缓解", cells)

    def test_default_compact_groups_reports_without_expanding_transcription(self):
        sources = [asset(f"r{i}") for i in range(9)]
        records = [analysis(f"r{i}", "laboratory", findings=[f"指标{i}：正常。", "其余指标未见明显异常。"])
                   for i in range(9)]
        for item in records:
            item["transcription"] = "不应反复铺陈的完整报告文字" * 100
        maker = AdvancedPPTMaker({}, sources, records)
        deck = Presentation(maker.build())
        self.assertLessEqual(len(deck.slides), 6)
        self.assertEqual({entry["asset_id"] for entry in maker.evidence_manifest}, {f"r{i}" for i in range(9)})
        combined = "\n".join(slide_text(slide) for slide in deck.slides)
        self.assertNotIn("完整报告文字", combined)
        self.assertFalse(any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for slide in deck.slides for shape in slide.shapes))
        self.assertFalse(any(shape.has_table for slide in deck.slides for shape in slide.shapes))
        for i in range(9):
            self.assertIn(f"指标{i}：正常。", combined)

    def test_explicit_plan_omits_unselected_images_and_unrelated_default_sections(self):
        maker = AdvancedPPTMaker({"baseline": {"key_exams": "不应自动追加"},
                                  "presentation_plan": {"slides": [{"type": "content", "title": "完整病例",
                                      "sections": [{"text": "已核对结果"}], "source_ids": ["used"]}]}},
                                 [asset("used"), asset("unused")], [analysis("used"), analysis("unused")])
        deck = Presentation(maker.build())
        self.assertEqual(len(deck.slides), 1)
        self.assertNotIn("不应自动追加", slide_text(deck.slides[0]))
        self.assertEqual([entry["asset_id"] for entry in maker.evidence_manifest], ["used"])
        self.assertEqual(maker.evidence_manifest[0]["kind"], "summary")
        self.assertIn("used", deck.slides[0].notes_slide.notes_text_frame.text)

    def test_table_overflow_retains_all_cells_and_native_table_shapes_fit(self):
        rows = [[f"日期{i}", "连续结果" * 30 + f"终点{i}"] for i in range(12)]
        maker = AdvancedPPTMaker({"presentation_plan": [{"type": "table", "title": "影像检查对照",
                                  "table": {"columns": ["日期", "影像结论"], "rows": rows,
                                            "column_widths": [1, 5]}, "source_ids": ["source"]}]},
                                 [asset("source")], [])
        deck = Presentation(maker.build())
        self.assertGreater(len(deck.slides), 1)
        combined = "".join(cell.text.replace("\n", "") for slide in deck.slides for shape in slide.shapes
                           if shape.has_table for row in shape.table.rows for cell in row.cells)
        for row in rows:
            for cell in row:
                self.assertIn(cell, combined)
        for slide in deck.slides:
            for shape in slide.shapes:
                self.assertGreaterEqual(shape.left, 0)
                self.assertGreaterEqual(shape.top, 0)
                self.assertLessEqual(shape.left + shape.width, deck.slide_width)
                self.assertLessEqual(shape.top + shape.height, deck.slide_height)

    def test_oversized_table_cell_splits_without_losing_any_text(self):
        content = "长篇结果" * 600 + "最终结论END"
        maker = AdvancedPPTMaker({"presentation_plan": [{"title": "详细检查",
                                  "table": {"columns": ["项目", "结果"], "rows": [["唯一项目", content]],
                                            "column_widths": [1, 5]}}]})
        deck = Presentation(maker.build())
        text = "".join(shape.table.cell(i, 1).text.replace("\n", "") for slide in deck.slides
                       for shape in slide.shapes if shape.has_table for i in range(1, len(shape.table.rows)))
        self.assertEqual(text, content)

    def test_invalid_planned_crop_never_falls_back_to_unselected_full_image(self):
        maker = AdvancedPPTMaker({"presentation_plan": [{"title": "核对结果", "source_ids": ["a"],
                                  "images": [{"asset_id": "a", "region_id": "invented"},
                                             {"asset_id": "not-present"}]}]}, [asset("a")], [analysis("a")])
        deck = Presentation(maker.build())
        self.assertFalse(any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for slide in deck.slides for shape in slide.shapes))
        self.assertTrue(maker.warnings)
        self.assertEqual(maker.evidence_manifest[0]["kind"], "summary")

    def test_missing_region_id_uses_only_an_existing_reviewed_crop(self):
        spec = {"title": "精选病理", "images": [{"asset_id": "a"}]}
        source = asset("a")
        maker = AdvancedPPTMaker({"presentation_plan": [spec]}, [source],
                                 [analysis("a", regions=[{"id": "reviewed", "box": [500, 0, 1000, 1000]}])])
        deck = Presentation(maker.build())
        picture = next(shape for shape in deck.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)
        with Image.open(io.BytesIO(picture.image.blob)) as image:
            self.assertEqual(image.size, (400, 400))
        self.assertEqual(maker.evidence_manifest[0]["region_ids"], ["reviewed"])
        no_regions = AdvancedPPTMaker({"presentation_plan": [spec]}, [source], [analysis("a")])
        no_regions_deck = Presentation(no_regions.build())
        self.assertFalse(any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE
                             for slide in no_regions_deck.slides for shape in slide.shapes))

    def test_default_summary_notes_keep_additional_findings_without_slide_repetition(self):
        findings = [f"完整观察{i}。" for i in range(7)]
        maker = AdvancedPPTMaker({}, [asset("a")], [analysis("a", "laboratory", findings=findings)])
        deck = Presentation(maker.build())
        notes = "\n".join(slide.notes_slide.notes_text_frame.text for slide in deck.slides)
        self.assertTrue(all(finding in notes for finding in findings))
        visible_text = "\n".join(slide_text(slide) for slide in deck.slides)
        self.assertIn(findings[2], visible_text)
        self.assertNotIn(findings[3], visible_text)

    def test_full_mode_preserves_original_pixels_without_a_crop(self):
        source = asset("a")
        maker = AdvancedPPTMaker({"presentation_plan": [{"title": "完整图像",
                                  "images": [{"asset_id": "a", "display_mode": "full"}]}]}, [source])
        deck = Presentation(maker.build())
        picture = next(shape for shape in deck.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)
        self.assertEqual(picture.image.blob, source["image_bytes"])
        self.assertAlmostEqual(picture.width / picture.height, 2.0, places=5)
        self.assertEqual(maker.evidence_manifest[0]["region_ids"], [])
        self.assertEqual(maker.evidence_manifest[0]["display_modes"], ["full"])
        self.assertEqual(maker.evidence_manifest[0]["kind"], "images")

    def test_full_mode_cannot_silently_override_a_bad_crop_reference(self):
        for mode in ("full", "unknown", None, []):
            maker = AdvancedPPTMaker({"presentation_plan": [{"title": "原始结果", "images": [
                {"asset_id": "a", "region_id": "invented", "display_mode": mode}]}]}, [asset("a")])
            deck = Presentation(maker.build())
            self.assertFalse(any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for slide in deck.slides for shape in slide.shapes))
            self.assertTrue(maker.warnings)

    def test_wide_image_sidebars_preserve_both_complete_images_and_editable_text(self):
        sources = [asset("portrait", width=400, height=800), asset("panel", width=800, height=600)]
        for layout in ("image_left", "image_right"):
            maker = AdvancedPPTMaker({"presentation_plan": [{"title": "病理检查", "layout": layout,
                "sections": [{"heading": "病理诊断", "text": "• 完整的病理诊断。\n• Ki-67约10%。"}],
                "images": [{"asset_id": source["id"], "display_mode": "full"} for source in sources]}]}, sources)
            deck = Presentation(maker.build())
            self.assertEqual(len(deck.slides), 1)
            pictures = [shape for shape in deck.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
            self.assertEqual(len(pictures), 2)
            for picture, source in zip(pictures, sources):
                self.assertEqual(picture.image.blob, source["image_bytes"])
                self.assertAlmostEqual(picture.width / picture.height, source["width"] / source["height"], places=5)
                self.assertLessEqual(picture.left + picture.width, deck.slide_width)
                self.assertLessEqual(picture.top + picture.height, deck.slide_height)
            bodies = [shape for shape in deck.slides[0].shapes if shape.has_text_frame and "Ki-67" in shape.text]
            self.assertTrue(bodies)
            if layout == "image_left":
                self.assertLess(max(picture.left + picture.width for picture in pictures), bodies[0].left)
            else:
                self.assertLess(bodies[0].left + bodies[0].width, min(picture.left for picture in pictures))

    def test_default_pathology_prefers_complete_group_and_does_not_make_a_table(self):
        records = [analysis("a", regions=[{"id": "detail", "label": "诊断文字", "box": [500, 0, 1000, 1000]},
                                           {"id": "whole", "label": "完整显微图组", "box": [0, 0, 1000, 1000]}])]
        maker = AdvancedPPTMaker({}, [asset("a")], records)
        deck = Presentation(maker.build())
        self.assertFalse(any(shape.has_table for slide in deck.slides for shape in slide.shapes))
        images = [entry for entry in maker.evidence_manifest if entry["kind"] == "images"]
        self.assertEqual(images[0]["region_ids"], ["whole"])
        self.assertIn("• 样本检测结果。", "\n".join(slide_text(slide) for slide in deck.slides))

    def test_report_pair_retains_complete_report_and_large_panel_above_editable_summary(self):
        sources = [asset("report", width=600, height=1000), asset("panel", width=1000, height=600)]
        summary = ("术后病理提示神经内分泌肿瘤，类型考虑不典型类癌。"
                   "免疫组化显示神经内分泌标志物阳性，增殖指数约百分之十。"
                   "再次活检考虑伴有神经内分泌分化的非小细胞肺癌，两次标本诊断分别保留。"
                   "两个标本的组织学结果和免疫表型分别列示。")
        self.assertTrue(90 < len(summary) < 131)
        spec = {"title": "病理及免疫组化", "layout": "report_pair", "sections": [{"text": summary}],
                "images": [{"asset_id": "report", "display_mode": "full", "label": "完整报告"},
                           {"asset_id": "panel", "display_mode": "full", "label": "完整显微图组"}]}
        maker = AdvancedPPTMaker({"presentation_theme": "standard_case", "presentation_plan": [spec]}, sources)
        deck = Presentation(maker.build())
        self.assertEqual(len(deck.slides), 1)
        pictures = [shape for shape in deck.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
        self.assertEqual(len(pictures), 2)
        for picture, source in zip(pictures, sources):
            self.assertEqual(picture.image.blob, source["image_bytes"])
            self.assertAlmostEqual(picture.width / picture.height, source["width"] / source["height"], places=5)
        body = [shape for shape in deck.slides[0].shapes if shape.has_text_frame
                and shape.top >= 4.5 * 914400 and shape.top < 7 * 914400]
        self.assertEqual("".join(shape.text for shape in body), summary)
        self.assertTrue(all(shape.text_frame.paragraphs[0].font.size.pt == 18 for shape in body))
        self.assertLess(pictures[0].left + pictures[0].width, pictures[1].left)
        self.assertLess(pictures[1].top + pictures[1].height, min(shape.top for shape in body))
        self.assertTrue(all(shape.top + shape.height < 6.81 * 914400 for shape in body))
        self.assertEqual([entry["asset_id"] for entry in maker.evidence_manifest], ["report", "panel"])

    def test_report_pair_never_drops_overflow_text_or_accepts_missing_pictures(self):
        text = "随访记录中的完整事实。" * 80 + "末尾结果保留。"
        spec = {"title": "检查结果", "layout": "report_pair", "sections": [{"text": text}],
                "images": [{"asset_id": "report", "display_mode": "full"},
                           {"asset_id": "panel", "display_mode": "full"}]}
        sources = [asset("report"), asset("panel")]
        deck = Presentation(AdvancedPPTMaker({"presentation_plan": [spec]}, sources).build())
        self.assertGreater(len(deck.slides), 1)
        body = "".join(shape.text for slide in deck.slides for shape in slide.shapes
                       if shape.has_text_frame and ("随访" in shape.text or "完整事实" in shape.text or "末尾" in shape.text))
        self.assertEqual(body, text)
        with self.assertRaises(ValueError):
            AdvancedPPTMaker({"presentation_plan": [spec]}, [sources[0]]).build()

    def test_default_serial_imaging_remains_an_editable_comparison_table(self):
        records = [analysis("a", "imaging", findings=["病灶20 mm。"]), analysis("b", "imaging", findings=["病灶18 mm。"])]
        records[1]["date"] = "2026-10-20"
        deck = Presentation(AdvancedPPTMaker({}, [asset("a"), asset("b")], records).build())
        tables = [shape.table for slide in deck.slides for shape in slide.shapes if shape.has_table]
        self.assertEqual(len(tables), 1)
        cells = "\n".join(cell.text for row in tables[0].rows for cell in row.cells)
        self.assertIn("病灶20 mm。", cells)
        self.assertIn("病灶18 mm。", cells)

    def test_standard_case_theme_uses_white_background_black_title_and_red_rule(self):
        maker = AdvancedPPTMaker({"presentation_theme": "standard_case", "presentation_plan": [
            {"type": "cover", "title": "病例汇报", "subtitle": "治疗经过"},
            {"type": "content", "title": "病例介绍", "sections": [{"text": "患者临床资料。"}]}]})
        deck = Presentation(maker.build())
        self.assertEqual(str(deck.slides[0].shapes[0].fill.fore_color.rgb), "FFFFFF")
        title = next(shape for shape in deck.slides[1].shapes if shape.has_text_frame and shape.text == "病例介绍")
        self.assertEqual(str(title.text_frame.paragraphs[0].font.color.rgb), "000000")
        self.assertEqual(title.text_frame.paragraphs[0].font.name, "Times New Roman")
        self.assertIn('typeface="宋体"', title._element.xml)
        self.assertTrue(any(shape.height < 50000 and str(shape.fill.fore_color.rgb) == "C00000"
                            for shape in deck.slides[1].shapes if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE))


class ClinicalTimelineTests(unittest.TestCase):
    def test_ten_nodes_are_editable_in_two_columns_and_plan_is_distinct(self):
        events = [{"date": f"2024-{i + 1:02d}", "phase": "治疗阶段",
                   "event_type": "Treatment", "event": f"方案{i}联合治疗"} for i in range(9)]
        events.append({"date": "2025-01", "event_type": "Treatment Plan", "event": "药物甲联合药物乙"})
        deck = Presentation(AdvancedPPTMaker({"presentation_theme": "standard_case", "presentation_plan": [
            {"type": "timeline", "title": "诊疗时间轴", "events": events}]}).build())
        self.assertEqual(len(deck.slides), 1)
        slide = deck.slides[0]
        self.assertTrue(all(e["date"] in slide_text(slide) and e["event"] in slide_text(slide) for e in events))
        self.assertIn("后续建议", slide_text(slide))
        self.assertFalse(any(s.has_table or s.shape_type == MSO_SHAPE_TYPE.PICTURE for s in slide.shapes))
        dates = [s for s in slide.shapes if s.has_text_frame and s.text in {e["date"] for e in events}]
        self.assertEqual(len({s.left for s in dates}), 2)
        self.assertTrue(all(s.text_frame.paragraphs[0].font.size.pt == 16 for s in dates))
        for shape in slide.shapes:
            self.assertGreaterEqual(shape.left, 0)
            self.assertGreaterEqual(shape.top, 0)
            self.assertLessEqual(shape.left + shape.width, deck.slide_width)
            self.assertLessEqual(shape.top + shape.height, deck.slide_height)

    def test_long_timeline_continues_without_losing_body_or_events(self):
        body = "病程原文内容完整保留。" * 70 + "末尾观察END"
        events = [{"date": "2024-01", "event_type": "Treatment", "event": body}]
        events += [{"date": f"2025-{i + 1:02d}", "event": f"事件编号{i:02d}"} for i in range(12)]
        deck = Presentation(AdvancedPPTMaker({"presentation_plan": [
            {"type": "timeline", "title": "诊疗时间轴", "events": events}]}).build())
        self.assertGreater(len(deck.slides), 1)
        text = "".join(shape.text.replace("\n", "").replace("\v", "")
                       for slide in deck.slides for shape in slide.shapes if shape.has_text_frame
                       and shape.text_frame.paragraphs[0].font.size == Pt(16)
                       and not shape.text_frame.paragraphs[0].font.bold)
        self.assertEqual(text.count("病程原文内容完整保留。"), 70)
        self.assertIn("末尾观察END", text)
        self.assertTrue(all(e["event"] in text for e in events[1:]))
        for slide in deck.slides:
            self.assertTrue(all(s.top + s.height <= deck.slide_height for s in slide.shapes))

    def test_compact_default_includes_grounded_timeline_before_treatment_details(self):
        data = {"timeline_events": [{"date": "2024-03", "phase": "治疗阶段", "event": "方案甲"}],
                "treatments": [{"phase": "后续治疗", "regimen": "方案甲"}]}
        deck = Presentation(AdvancedPPTMaker(data).build())
        text = [slide_text(s) for s in deck.slides]
        timeline = next(i for i, t in enumerate(text) if "诊疗时间轴" in t)
        details = next(i for i, t in enumerate(text) if "系统治疗经过" in t)
        self.assertLess(timeline, details)
        self.assertIn("方案甲", text[timeline])

    def test_timeline_rejects_missing_dates_and_incompatible_content(self):
        for spec in ({"events": []}, {"events": [{"event": "缺少日期"}]},
                     {"events": [{"date": "2025", "event": "事件"}], "bullets": ["不能静默丢弃"]}):
            with self.assertRaises(ValueError):
                AdvancedPPTMaker({"presentation_plan": [{"type": "timeline", **spec}]}).build()


if __name__ == "__main__":
    unittest.main()
