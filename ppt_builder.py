"""Editable case slides and high-resolution clinical evidence placement."""

import copy
import io
import math
import re
import unicodedata

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from vision_pipeline import crop_region


PPT_LAYOUT_VERSION = "2026-10-22.5"

# Keep short clinical names and Latin assay identifiers intact at manual line breaks.
_CLINICAL_WORDS = (
    "纳武利尤单抗", "卡度尼利单抗", "卡度尼利", "替莫唑胺", "卡培他滨", "索凡替尼",
    "白蛋白紫杉醇", "白蛋白结合型", "紫杉醇", "神经内分泌", "免疫组化",
    "非小细胞肺癌", "不典型类癌", "肿瘤标志物", "楔形切除术", "电磁导航", "胸腔镜",
)
_WRAP_TOKENS = re.compile(
    "|".join(re.escape(word) for word in sorted(_CLINICAL_WORDS, key=len, reverse=True))
    + r"|[A-Za-z0-9]+(?:[._/+−-][A-Za-z0-9]+)*(?:\([^()\s]*\))*|."
)
_LINE_START_FORBIDDEN = frozenset("，。；：！？、）】》〉〕］｝」』”’％,.;:!?)]}%")


class AdvancedPPTMaker:
    CATEGORY_NAMES = {
        "imaging": "影像学检查", "pathology": "病理与免疫组化",
        "genetic": "基因检测", "genetics": "基因检测", "genomic": "基因检测",
        "laboratory": "检验报告", "lab": "检验报告", "medical_record": "病历资料",
        "clinical": "病历资料", "record": "病历资料", "other": "检查资料", "unknown": "检查资料",
    }

    def __init__(self, data, assets=None, analyses=None):
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)
        self.data = self.clean_data(data)
        self.C_PRI = RGBColor(115, 21, 40)
        self.C_ACC = RGBColor(0, 51, 102)
        self.standard_case_theme = self.data.get("presentation_theme") == "standard_case"
        if self.standard_case_theme:
            self.C_PRI = RGBColor(192, 0, 0)
            self.C_ACC = RGBColor(40, 40, 40)
        self.assets = {asset["id"]: asset for asset in (assets or [])
                       if isinstance(asset, dict) and isinstance(asset.get("id"), str)
                       and asset["id"] and asset.get("image_bytes")}
        self.analyses = {item["asset_id"]: copy.deepcopy(item) for item in (analyses or [])
                         if isinstance(item, dict) and isinstance(item.get("asset_id"), str)
                         and item["asset_id"] in self.assets}
        self.source_labels = {}
        for i, (asset_id, asset) in enumerate(self.assets.items(), 1):
            source_index = asset.get("source_index")
            if type(source_index) is not int or source_index <= 0:
                source_index = i
            self.source_labels[asset_id] = f"资料 {source_index:02d}"
        self.warnings = []
        self.evidence_manifest = []
        self.evidence = self._resolve_evidence()

    def clean_data(self, data):
        # Formatting must never infer treatment phases or mutate the caller's data.
        return copy.deepcopy(data)

    def _resolve_evidence(self):
        """Resolve source IDs once, with explicit fallbacks for incomplete model links."""
        resolved, linked = [], set()
        excluded = {item for item in (self.data.get("excluded_asset_ids") or []) if isinstance(item, str)}
        treatment_count = len(self.data.get("treatments") or [])
        for link in self.data.get("evidence_links") or []:
            if not isinstance(link, dict):
                continue
            asset_id = link.get("asset_id")
            if not isinstance(asset_id, str):
                self.warnings.append("忽略了无效的图片引用。")
                continue
            if asset_id in excluded:
                continue
            if asset_id not in self.assets:
                self.warnings.append("忽略了不存在的图片引用。")
                continue
            section, tx_index = link.get("section", "appendix"), link.get("treatment_index")
            if not isinstance(section, str) or section not in {"baseline", "treatment", "current_admission", "appendix"}:
                section = "appendix"
            if section == "treatment" and (
                not isinstance(tx_index, int) or isinstance(tx_index, bool)
                or not 0 <= tx_index < treatment_count
            ):
                section, tx_index = "appendix", None
            if section == "current_admission" and not self.data.get("current_admission"):
                section = "appendix"
            if section != "treatment":
                tx_index = None
            analysis = self.analyses.get(asset_id, {})
            regions = [region for region in (analysis.get("regions") or []) if isinstance(region, dict)]
            requested = link.get("region_ids") or []
            if not isinstance(requested, list):
                requested = [requested]
            known = {region.get("id") for region in regions if isinstance(region.get("id"), str)}
            if requested and any(not isinstance(region_id, str) or region_id not in known for region_id in requested):
                # A hallucinated region is not safe to interpret as another crop.
                regions, section, tx_index = [], "appendix", None
                self.warnings.append(f"{self.source_labels[asset_id]} 的区域引用无效，已保留完整图片。")
            elif requested:
                regions = [region for region in regions if region.get("id") in requested]
            resolved.append({"asset_id": asset_id, "section": section,
                             "treatment_index": tx_index, "caption": self._as_text(link.get("caption")),
                             "regions": regions})
            linked.add(asset_id)
        for asset_id in self.assets:
            if asset_id not in linked and asset_id not in excluded:
                regions = self.analyses.get(asset_id, {}).get("regions") or []
                resolved.append({"asset_id": asset_id, "section": "appendix", "treatment_index": None,
                                 "caption": "", "regions": regions})
        return resolved

    @staticmethod
    def _as_text(value):
        if value is None:
            return ""
        if isinstance(value, list):
            return "\n".join(str(item) for item in value if item is not None)
        return str(value)

    @staticmethod
    def _wrap_lines(text, max_units=80):
        """CJK-aware wrapping that preserves medical names and every source character."""
        lines = []
        for original in str(text).splitlines():
            line, units = "", 0
            for token in _WRAP_TOKENS.findall(original):
                cost = sum(2 if unicodedata.east_asian_width(char) in "WF" else 1 for char in token)
                # Only an individual token wider than the whole box needs splitting.
                pieces = list(token) if cost > max_units else [token]
                for piece in pieces:
                    piece_cost = sum(2 if unicodedata.east_asian_width(char) in "WF" else 1 for char in piece)
                    if line and units + piece_cost > max_units:
                        # Chinese closing punctuation must not become an isolated next
                        # line. Move the preceding complete token with it, retaining
                        # clinical identifiers and every original character.
                        previous_tokens = _WRAP_TOKENS.findall(line)
                        if piece[0] in _LINE_START_FORBIDDEN:
                            if len(previous_tokens) > 1:
                                moved = previous_tokens.pop()
                                lines.append("".join(previous_tokens))
                                line = moved
                                units = sum(2 if unicodedata.east_asian_width(char) in "WF" else 1
                                            for char in moved)
                            # A single unbreakable token keeps its closing mark.
                        else:
                            lines.append(line)
                            line, units = "", 0
                    line += piece
                    units += piece_cost
            lines.append(line)
        return lines or [""]

    @staticmethod
    def _phase_label(value):
        labels = {"baseline": "基线评估", "treatment": "治疗阶段",
                  "current_admission": "本次评估", "appendix": "补充检查",
                  "evaluation": "疗效评估", "follow_up": "随访", "followup": "随访",
                  "plan": "治疗计划", "surgery": "手术", "adjuvant": "辅助治疗",
                  "neoadjuvant": "新辅助治疗"}
        value = str(value or "")
        return labels.get(value.strip().lower(), value)

    def _apply_font(self, font):
        font.name = "Times New Roman" if self.standard_case_theme else "Microsoft YaHei"
        if self.standard_case_theme:
            east_asian = font._element.find(qn("a:ea"))
            if east_asian is None:
                east_asian = OxmlElement("a:ea")
                font._element.append(east_asian)
            east_asian.set("typeface", "宋体")

    def _textbox(self, slide, text, left, top, width, height, size=18, color=None, bold=False):
        box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_top = tf.margin_bottom = 0
        tf.margin_left = tf.margin_right = 0
        paragraph = tf.paragraphs[0]
        paragraph.text = text
        self._apply_font(paragraph.font)
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = color or RGBColor(40, 40, 40)
        paragraph.line_spacing = 1.1
        return box

    def _make_body_pages(self, title, sections):
        """Paginate editable body text before placing it, without shrinking or dropping facts.

        Each section is (heading, text, body_color, body_bold). Conservative CJK
        widths and 1.5 font-size line boxes leave room for actual Office font metrics.
        """
        sections = list(sections)
        if len(self._wrap_lines(title, 68)) > 2:
            sections.insert(0, ("", title, self.C_PRI, True))
            title = title.split("：", 1)[0] if "：" in title else "病例资料"
        blocks = []
        for heading, text, color, bold in sections:
            rows = []
            if heading:
                rows.extend((line, 20, self.C_PRI, True) for line in self._wrap_lines(heading, 76))
            body = self._as_text(text)
            if body:
                rows.extend((line, 18, color, bold) for line in self._wrap_lines(body, 84))
            if rows:
                blocks.append(rows)
        pages, current, height = [], [], 0.0
        for block in blocks:
            # Keep a heading and its first body line together at page boundaries.
            initial_height = sum(row[1] * 1.5 / 72 for row in block[:2])
            if current and height + initial_height > 5.75:
                pages.append(current)
                current, height = [], 0.0
            for row in block:
                row_height = row[1] * 1.5 / 72
                if current and height + row_height > 5.75:
                    pages.append(current)
                    current, height = [], 0.0
                current.append((row, height, row_height))
                height += row_height
            height += .12
        if current or not pages:
            pages.append(current)
        for page_index, rows in enumerate(pages):
            slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
            page_title = title if len(pages) == 1 else f"{title}（{page_index + 1}/{len(pages)}）"
            self.add_header(slide, page_title)
            for (text, size, color, bold), offset, row_height in rows:
                self._textbox(slide, text, .8, 1.2 + offset, 11.7, row_height,
                              size=size, color=color, bold=bold)

    def _evidence_footer(self, slide, asset_id, date):
        text = self.source_labels[asset_id]
        if date:
            text += f"　{date}"
        self._textbox(slide, text, .8, 7.02, 11.7, .25, size=12, color=RGBColor(90, 90, 90))

    def _record_evidence_slide(self, asset_id, link, region_ids, kind, display_modes=None):
        entry = {
            "slide_number": len(self.prs.slides), "asset_id": asset_id,
            "source_label": self.source_labels[asset_id], "region_ids": region_ids,
            "section": link["section"], "treatment_index": link["treatment_index"], "kind": kind,
        }
        if display_modes is not None:
            entry["display_modes"] = list(display_modes)
        self.evidence_manifest.append(entry)

    @staticmethod
    def _sentence_units(text):
        """Keep complete source sentences/items; never invent a shorter medical claim."""
        units = []
        for paragraph in str(text).splitlines():
            # Decimal points and gene/protein notation are deliberately not delimiters.
            units.extend(piece.strip() for piece in re.split(r"(?<=[。！？!?])", paragraph)
                         if piece.strip())
        return units

    def _take_complete_units(self, units, max_units, max_lines):
        selected, count = [], 0
        while units:
            unit_lines = self._wrap_lines(units[0], max_units)
            if count + len(unit_lines) > max_lines:
                break
            selected.append(units.pop(0))
            count += len(unit_lines)
        return "\n".join("\n".join(self._wrap_lines(unit, max_units)) for unit in selected)

    def _detail_pages(self, units, max_units=80, max_lines=12):
        """Prefer sentence boundaries between pages, splitting only oversized units."""
        units = list(units)
        pages = []
        while units:
            text = self._take_complete_units(units, max_units, max_lines)
            if text:
                pages.append(text)
                continue
            oversized = units.pop(0)
            clauses = [piece for piece in re.split(r"(?<=[；;，,])", oversized) if piece]
            if len(clauses) > 1:
                # A long immunohistochemistry list can span pages at marker boundaries.
                units[0:0] = clauses
                continue
            lines = self._wrap_lines(oversized, max_units)
            pages.extend("\n".join(lines[offset:offset + max_lines])
                         for offset in range(0, len(lines), max_lines))
        return pages

    def _evidence_batches(self, section, treatment_index):
        links = [link for link in self.evidence
                 if link["section"] == section and link["treatment_index"] == treatment_index]
        by_id = {link["asset_id"]: link for link in links}
        groups = {}
        for group in self.data.get("evidence_groups") or []:
            if not isinstance(group, dict) or group.get("reason") != "matching_report_content":
                continue
            asset_ids = group.get("asset_ids")
            if not isinstance(asset_ids, list):
                continue
            members = list(dict.fromkeys(aid for aid in asset_ids if isinstance(aid, str) and aid in by_id))
            if len(members) > 1:
                for aid in members:
                    groups[aid] = members
        handled = set()
        for link in links:
            aid = link["asset_id"]
            if aid in handled:
                continue
            members = [candidate for candidate in groups.get(aid, [aid]) if candidate not in handled]
            handled.update(members)
            yield [by_id[candidate] for candidate in members]

    def _evidence_page_footer(self, slide, asset_ids):
        labels = []
        for aid in dict.fromkeys(asset_ids):
            date = self._as_text(self.analyses.get(aid, {}).get("date"))
            labels.append(self.source_labels[aid] + (f"　{date}" if date else ""))
        self._textbox(slide, "　 /　 ".join(labels), .8, 7.02, 11.7, .25, size=12,
                      color=RGBColor(90, 90, 90))

    def _place_evidence_picture(self, slide, picture_data, left, top, width, height, ordinal):
        scale = min(width / picture_data["width"], height / picture_data["height"])
        placed_width = picture_data["width"] * scale
        placed_height = picture_data["height"] * scale
        picture = slide.shapes.add_picture(io.BytesIO(picture_data["bytes"]),
            Inches(left + (width - placed_width) / 2), Inches(top),
            width=Inches(placed_width), height=Inches(placed_height))
        picture.name = f"{self.source_labels[picture_data['asset_id']]}_图 {ordinal}"
        return placed_height

    def _plan_source_ids(self, spec):
        """Resolve only existing source IDs; display labels never expose upload names."""
        requested = list(spec.get("source_ids") or [])
        requested.extend(item.get("asset_id") for item in spec.get("images") or []
                         if isinstance(item, dict))
        excluded = self.data.get("excluded_asset_ids") or []
        resolved = []
        for aid in requested:
            if not isinstance(aid, str) or aid not in self.assets or aid in excluded:
                self.warnings.append("汇报页中存在无效或已排除的资料引用，已忽略。")
            elif aid not in resolved:
                resolved.append(aid)
        return resolved

    def _plan_pictures(self, spec):
        pictures = []
        for item in spec.get("images") or []:
            if not isinstance(item, dict):
                continue
            aid, region_id = item.get("asset_id"), item.get("region_id")
            if not isinstance(aid, str) or aid not in self.assets or aid in (self.data.get("excluded_asset_ids") or []):
                continue
            asset = self.assets[aid]
            regions = self.analyses.get(aid, {}).get("regions") or []
            mode = item.get("display_mode", "region")
            if mode not in ("region", "full") or (mode == "full" and "region_id" in item):
                self.warnings.append(f"{self.source_labels[aid]} 的图片显示方式无效，已忽略该图。")
                continue
            region = None
            if mode == "region":
                if "region_id" in item:
                    region = next((r for r in regions if isinstance(r, dict) and r.get("id") == region_id), None)
                else:
                    region = next((r for r in regions if isinstance(r, dict) and r.get("id")), None)
                    region_id = region.get("id") if region else None
                if region is None:
                    self.warnings.append(f"{self.source_labels[aid]} 的精选区域不存在，已忽略该图。")
                    continue
            try:
                raw = asset["image_bytes"] if mode == "full" else crop_region(asset, region)
                with Image.open(io.BytesIO(raw)) as image:
                    width, height = image.size
            except (ValueError, TypeError, KeyError, OSError):
                self.warnings.append(f"{self.source_labels[aid]} 的精选区域无法读取，已忽略该图。")
                continue
            label = self._as_text(item.get("label") or (region or {}).get("label"))
            pictures.append({"bytes": raw, "asset_id": aid, "region_id": region_id, "display_mode": mode,
                             "label": label, "width": width, "height": height,
                             "aspect": width / height})
        return pictures

    def _compact_source_footer(self, slide, source_ids, spec, pictures):
        if source_ids:
            # Numerous sources remain in the notes and manifest; use compact display labels.
            label = "　".join(self.source_labels[aid] for aid in source_ids)
            lines = self._wrap_lines(label, 170)
            self._textbox(slide, "\n".join(lines[:2]), .65, 7.01, 11.7, .38, 10,
                          RGBColor(100, 100, 100))
        section = spec.get("section", "appendix")
        if section not in ("baseline", "treatment", "current_admission", "appendix"):
            section = "appendix"
        tx_index = spec.get("treatment_index") if section == "treatment" else None
        for aid in source_ids:
            selected = [p["region_id"] for p in pictures if p["asset_id"] == aid and p["region_id"]]
            modes = [p["display_mode"] for p in pictures if p["asset_id"] == aid]
            self._record_evidence_slide(aid, {"section": section, "treatment_index": tx_index},
                                        selected, "images" if selected or any(p["asset_id"] == aid for p in pictures)
                                        else "summary", display_modes=modes)
        slide.notes_slide.notes_text_frame.text = "\n".join(
            ["资料来源（用于追溯）"] + [f"{self.source_labels[aid]} | {aid}" for aid in source_ids]
            + ([self._as_text(spec.get("notes"))] if spec.get("notes") else []))

    def _compact_text_rows(self, spec, width):
        rows = []
        if spec.get("subtitle"):
            rows.extend((line, 19, self.C_ACC, True) for line in
                        self._wrap_lines(self._as_text(spec["subtitle"]), int(width * 72 / 10.5)))
        sections = spec.get("sections") or []
        if spec.get("bullets"):
            sections = list(sections) + [{"text": spec["bullets"]}]
        for section in sections:
            if not isinstance(section, dict):
                continue
            heading = self._as_text(section.get("heading"))
            if heading:
                rows.extend((line, 19, self.C_PRI, True) for line in
                            self._wrap_lines(heading, int(width * 72 / 10.5)))
            text = self._as_text(section.get("text", section.get("body", "")))
            if text:
                rows.extend((line, 18, None, False) for line in
                            self._wrap_lines(text, int(width * 72 / 10)))
            if heading or text:
                rows.append(("", 5, None, False))
        return rows

    def _compact_table_rows(self, spec, width, max_height):
        table = spec.get("table")
        if not isinstance(table, dict) or not table.get("columns"):
            return None
        headers = [self._as_text(value) for value in table["columns"]]
        columns = len(headers)
        if columns > 8:
            raise ValueError("一页表格最多支持 8 列，请将表格按主题拆分。")
        proportions = table.get("column_widths") or [1] * columns
        if (not isinstance(proportions, list) or len(proportions) != columns or
                any(not isinstance(v, (int, float)) or isinstance(v, bool) or v <= 0 for v in proportions)):
            proportions = [1] * columns
        widths = [width * value / sum(proportions) for value in proportions]
        font_size = max(14, min(20, table.get("font_size", 17)))
        limits = [max(6, int((w - .18) * 72 / (font_size * .56))) for w in widths]
        line_height = font_size * 1.35 / 72
        header_lines = [self._wrap_lines(value, limits[i]) for i, value in enumerate(headers)]
        header_height = max(.43, max(map(len, header_lines)) * line_height + .15)
        max_lines = max(1, int((max_height - header_height - .2) / line_height))
        rows = []
        for raw_row in table.get("rows") or []:
            if not isinstance(raw_row, (list, tuple)):
                continue
            cells = [self._as_text(raw_row[i]) if i < len(raw_row) else "" for i in range(columns)]
            wrapped = [self._wrap_lines(value, limits[i]) for i, value in enumerate(cells)]
            # Preserve every cell, even if one row alone is longer than a slide.
            for start in range(0, max(map(len, wrapped)), max_lines):
                page_cells = ["\n".join(lines[start:start + max_lines]) for lines in wrapped]
                height = max(.44, max(len(lines[start:start + max_lines]) for lines in wrapped) * line_height + .16)
                rows.append((page_cells, height))
        return {"headers": ["\n".join(lines) for lines in header_lines], "header_height": header_height,
                "widths": widths, "rows": rows, "font_size": font_size}

    def _draw_compact_table(self, slide, table_data, rows, left, top):
        height = table_data["header_height"] + sum(row[1] for row in rows)
        frame = slide.shapes.add_table(len(rows) + 1, len(table_data["headers"]),
                                      Inches(left), Inches(top), Inches(sum(table_data["widths"])), Inches(height))
        table = frame.table
        for i, width in enumerate(table_data["widths"]):
            table.columns[i].width = Inches(width)
        for row_index, (values, row_height) in enumerate(
                [(table_data["headers"], table_data["header_height"])] + rows):
            table.rows[row_index].height = Inches(row_height)
            for col_index, value in enumerate(values):
                cell = table.cell(row_index, col_index)
                cell.text = value
                cell.margin_left = cell.margin_right = Inches(.09)
                cell.margin_top = cell.margin_bottom = Inches(.06)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                cell.fill.solid()
                cell.fill.fore_color.rgb = self.C_PRI if row_index == 0 else RGBColor(
                    *( (248, 241, 243) if row_index % 2 else (255, 255, 255)))
                for paragraph in cell.text_frame.paragraphs:
                    self._apply_font(paragraph.font)
                    paragraph.font.size = Pt(table_data["font_size"])
                    paragraph.font.bold = row_index == 0
                    paragraph.font.color.rgb = RGBColor(255, 255, 255) if row_index == 0 else RGBColor(40, 40, 40)
                    paragraph.line_spacing = 1.0

    def _compact_picture_grid(self, slide, pictures, left, top, width, height):
        if not pictures:
            return
        gap = .22
        panel_width = (width - gap * (len(pictures) - 1)) / len(pictures)
        for i, picture in enumerate(pictures):
            x = left + i * (panel_width + gap)
            label = picture["label"]
            if label:
                label_lines = self._wrap_lines(label, max(12, int(panel_width * 72 / 8)))
                label_height = max(.35, len(label_lines) * .23)
                self._textbox(slide, "\n".join(label_lines), x, top, panel_width, label_height,
                              13, self.C_ACC, True)
            else:
                label_height = .08
            self._place_evidence_picture(slide, picture, x, top + label_height,
                                         panel_width, max(.1, height - label_height), i + 1)

    def make_planned_slide(self, spec):
        """Render one compact editorial spec with editable text/table and selected original crops.

        Source references are explicit, and an explicit plan never automatically appends
        all reports or transcriptions. Overflow is paginated without cutting facts.
        """
        if not isinstance(spec, dict):
            raise ValueError("幻灯计划必须是对象。")
        if spec.get("type") == "timeline":
            self.make_timeline_spec(spec)
            return
        if spec.get("type") == "cover":
            previous = self.data.get("cover")
            self.data["cover"] = {"title": spec.get("title") or "病例汇报"}
            self.make_cover()
            if previous is None:
                self.data.pop("cover", None)
            else:
                self.data["cover"] = previous
            if spec.get("subtitle"):
                self._textbox(self.prs.slides[-1], self._as_text(spec["subtitle"]), 1.5, 6.15, 10, .7,
                              20, RGBColor(40, 40, 40) if self.standard_case_theme else RGBColor(255, 255, 255))
            self._compact_source_footer(self.prs.slides[-1], self._plan_source_ids(spec), spec, [])
            return
        pictures = self._plan_pictures(spec)
        report_pair = spec.get("layout") == "report_pair"
        if report_pair and len(pictures) != 2:
            raise ValueError("report_pair版式需要两张有效图片，依次为完整报告和完整图组。")
        sources = self._plan_source_ids(spec)
        title = self._as_text(spec.get("title") or "病例资料")
        # More than three explicit selections continue on another image page, never disappear.
        picture_batches = [pictures[i:i + 3] for i in range(0, len(pictures), 3)] or [[]]
        first = picture_batches[0]
        has_content = bool(spec.get("sections") or spec.get("bullets") or spec.get("table") or spec.get("subtitle"))
        table_columns = len((spec.get("table") or {}).get("columns") or [])
        image_sidebar = bool(first) and has_content and spec.get("layout") in ("image_left", "image_right")
        side_by_side = image_sidebar or (len(first) == 1 and has_content and table_columns <= 3
                                         and spec.get("layout") != "bottom")
        body_width = 6.5 if report_pair else (5.0 if image_sidebar else (7.05 if side_by_side else 12.0))
        body_left = 6.15 if report_pair else (7.65 if image_sidebar and spec.get("layout") == "image_left" else .65)
        body_height = 2.25 if report_pair else (5.65 if side_by_side or not first else (2.65 if has_content else 0))
        text_rows = self._compact_text_rows(spec, body_width)
        if report_pair:
            # Six readable 18pt lines fit below the complete microscopy panel.
            while text_rows and not text_rows[-1][0]:
                text_rows.pop()
        table = self._compact_table_rows(spec, body_width, body_height or 5.65)
        table_rows = list(table["rows"]) if table else []
        text_rows = list(text_rows)
        pages = []
        while text_rows or table_rows or not pages:
            is_first = not pages
            capacity = body_height if is_first else 5.65
            page_width = body_width
            selected_text, used = [], 0.0
            while text_rows:
                row = text_rows[0]
                row_height = row[1] * 1.45 / 72
                if used + row_height > capacity:
                    break
                # Keep a heading together with the next body line when possible.
                if row[3] and len(text_rows) > 1 and used + row_height + text_rows[1][1] * 1.45 / 72 > capacity:
                    break
                selected_text.append(text_rows.pop(0))
                used += row_height
            selected_table = []
            if not text_rows and table_rows:
                if selected_text:
                    used += .1
                header_h = table["header_height"]
                while table_rows and used + header_h + table_rows[0][1] <= capacity + .001:
                    selected_table.append(table_rows.pop(0))
                    used += selected_table[-1][1]
            if not selected_text and not selected_table and (text_rows or table_rows):
                # A first-page image strip can leave too little room for a table row.
                if is_first and first:
                    pages.append(([], [], True))
                    continue
                raise ValueError("本页内容无法在可读字号下排入，请缩短表头或拆分段落。")
            pages.append((selected_text, selected_table, is_first))
        total = len(pages) + len(picture_batches) - 1
        for page_index, (rows, selected_table, is_first) in enumerate(pages):
            slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
            self.add_header(slide, title if total == 1 else f"{title}（{page_index + 1}/{total}）")
            top = 4.52 if report_pair and is_first else 1.14
            for text, size, color, bold in rows:
                row_height = size * 1.45 / 72
                self._textbox(slide, text, body_left, top, body_width, row_height, size, color, bold)
                top += row_height
            if selected_table:
                self._draw_compact_table(slide, table, selected_table, body_left, top + (.1 if rows else 0))
            actual_pictures = first if is_first else []
            if actual_pictures:
                if report_pair:
                    self._compact_picture_grid(slide, actual_pictures[:1], .65, 1.14, 5.0, 5.65)
                    self._compact_picture_grid(slide, actual_pictures[1:], 6.15, 1.14, 6.5, 3.2)
                elif image_sidebar:
                    image_left = .65 if spec.get("layout") == "image_left" else 6.05
                    self._compact_picture_grid(slide, actual_pictures, image_left, 1.14, 6.6, 5.65)
                elif side_by_side:
                    self._compact_picture_grid(slide, actual_pictures, 7.95, 1.14, 4.7, 5.65)
                else:
                    image_top = 4.02 if has_content else 1.15
                    self._compact_picture_grid(slide, actual_pictures, .65, image_top, 12.0, 6.8 - image_top)
            self._compact_source_footer(slide, sources, spec, actual_pictures)
        for batch_index, batch in enumerate(picture_batches[1:], len(pages) + 1):
            slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
            self.add_header(slide, f"{title}（{batch_index}/{total}）")
            self._compact_picture_grid(slide, batch, .65, 1.15, 12.0, 5.65)
            self._compact_source_footer(slide, sources, spec, batch)

    def _compact_evidence_specs(self, section, treatment_index=None):
        grouped = {}
        for link in self.evidence:
            if link["section"] == section and link["treatment_index"] == treatment_index:
                category = self.analyses.get(link["asset_id"], {}).get("category", "other")
                title = self.CATEGORY_NAMES.get(category, "检查资料")
                grouped.setdefault(title, []).append(link)
        specs = []
        for title, links in grouped.items():
            rows, sections, selected, notes = [], [], [], []
            for link in links:
                aid = link["asset_id"]
                analysis = self.analyses.get(aid, {})
                findings = analysis.get("findings") or []
                if isinstance(findings, str):
                    findings = [findings]
                units = [unit for item in findings for unit in self._sentence_units(item)]
                concise = "\n".join(units[:3]) or link["caption"] or self._as_text(analysis.get("title"))
                rows.append([self._as_text(analysis.get("date")), concise, self.source_labels[aid]])
                heading = "　".join(value for value in [self._as_text(analysis.get("date")),
                                    self._as_text(analysis.get("title")) or self.source_labels[aid]] if value)
                sections.append({"heading": heading, "text": "\n".join("• " + unit for unit in units[:3]) or concise})
                notes.append(self.source_labels[aid] + "\n" + "\n".join(units))
                # Preserve one complete report body/image group instead of fragmenting it.
                if not units or (len(selected) < 2 and title in ("病理与免疫组化", "影像学检查", "基因检测")):
                    regions = [region for region in link["regions"] if isinstance(region, dict) and region.get("id")]
                    complete = [region for region in regions if re.search(
                        r"完整.*(?:报告|图|影像)|报告主体|图组|全貌", self._as_text(region.get("label")))]
                    region = next(iter(complete or regions), None)
                    if region:
                        selected.append({"asset_id": aid, "region_id": region["id"],
                                         "label": self.source_labels[aid] + "　" + self._as_text(region.get("label"))})
            # Dated serial measurements benefit from a table; a report's conclusions are prose.
            dated_comparison = title in ("影像学检查", "检验报告") and len({row[0] for row in rows if row[0]}) > 1
            spec = {"type": "table" if dated_comparison else "content", "title": title, "section": section,
                    "treatment_index": treatment_index, "source_ids": [link["asset_id"] for link in links],
                    "images": selected, "notes": "\n\n".join(notes)}
            if dated_comparison:
                spec["table"] = {"columns": ["日期", "主要结果", "来源"], "rows": rows,
                                 "column_widths": [1.3, 6.6, 1]}
            else:
                spec["sections"] = sections
                if len(selected) == 1:
                    spec["layout"] = "image_right"
            specs.append(spec)
        return specs

    def compact_plan(self):
        """A conservative grouped fallback for existing case JSON; no inferred line numbers."""
        base = self.data.get("baseline") or {}
        specs = [{"type": "cover", "title": (self.data.get("cover") or {}).get("title", "病例汇报")},
                 {"type": "content", "title": "病例概况", "section": "baseline",
                  "sections": [{"heading": heading, "text": base.get(key, "")} for heading, key in
                               (("患者信息", "patient_info"), ("主诉", "chief_complaint"),
                                ("临床诊断", "diagnosis"), ("关键检查", "key_exams"))]}]
        specs.extend(self._compact_evidence_specs("baseline"))
        if self.data.get("timeline_events"):
            specs.append({"type": "timeline", "title": "诊疗时间轴", "section": "treatment",
                          "events": self.data["timeline_events"]})
        treatments = self.data.get("treatments") or []
        if treatments:
            def label(tx):
                explicit = tx.get("line_label") or tx.get("line")
                return (f"第{explicit}线" if type(explicit) is int else self._as_text(explicit)) or self._phase_label(tx.get("phase", "治疗阶段"))
            specs.append({"type": "table", "title": "系统治疗经过", "section": "treatment",
                          "table": {"columns": ["治疗线数 / 阶段", "时间", "方案", "疗效与转归"],
                                    "column_widths": [1.5, 1.7, 3.4, 3.4],
                                    "rows": [[label(tx), tx.get("duration", ""), tx.get("regimen", ""),
                                              tx.get("response") or tx.get("imaging", "")] for tx in treatments]}})
            for index, tx in enumerate(treatments):
                specs.append({"type": "content", "title": label(tx), "section": "treatment", "treatment_index": index,
                              "sections": [{"heading": heading, "text": tx.get(key, "")} for heading, key in
                                           (("治疗时间", "duration"), ("治疗方案", "regimen"),
                                            ("影像学评估", "imaging"), ("肿瘤标志物", "markers"))]})
                specs.extend(self._compact_evidence_specs("treatment", index))
        current = self.data.get("current_admission") or {}
        if current:
            specs.append({"type": "content", "title": "当前评估与后续计划", "section": "current_admission",
                          "sections": [{"heading": heading, "text": current.get(key, "")} for heading, key in
                                       (("检查结果", "exams"), ("影像学评估", "imaging"), ("后续计划", "plan"))]})
        specs.extend(self._compact_evidence_specs("current_admission"))
        specs.extend(self._compact_evidence_specs("appendix"))
        summary = self.data.get("summary") or {}
        specs.append({"type": "content", "title": "病例总结", "sections": [
            {"text": summary if isinstance(summary, list) else summary.get("highlights", [])},
            {"heading": "病例思考" if isinstance(summary, dict) and summary.get("discussion") else "",
             "text": summary.get("discussion", []) if isinstance(summary, dict) else ""}]})
        return specs

    def make_evidence(self, section, treatment_index=None):
        """Fit complete clinical statements around uncropped, proportional evidence images."""
        for batch in self._evidence_batches(section, treatment_index):
            links_by_id = {link["asset_id"]: link for link in batch}
            units, seen_text, source_pictures = [], set(), []
            categories = []
            for link in batch:
                asset_id = link["asset_id"]
                asset = self.assets[asset_id]
                analysis = self.analyses.get(asset_id, {})
                category_value = analysis.get("category")
                categories.append(self.CATEGORY_NAMES.get(category_value, "检查资料")
                                  if isinstance(category_value, str) else "检查资料")
                report_title = self._as_text(analysis.get("title")).strip()
                caption = link["caption"].strip()
                parts = [report_title if caption and caption in report_title else (caption or report_title)]
                if report_title and report_title not in caption and caption not in report_title:
                    parts.append(report_title)
                parts.append(self._as_text(analysis.get("findings")))
                if not any(parts):
                    parts.append(self._as_text(analysis.get("transcription")))
                source_fingerprints = set()
                for part in parts:
                    for sentence in self._sentence_units(part):
                        fingerprint = re.sub(r"\s+", "", unicodedata.normalize("NFKC", sentence))
                        if fingerprint not in seen_text:
                            units.append(sentence)
                            source_fingerprints.add(fingerprint)
                # Deduplicate only across explicitly grouped copies of one report.
                # Repeated statements inside a source may belong to different dates.
                seen_text.update(source_fingerprints)
                pictures = []
                for region in link["regions"]:
                    try:
                        cropped = crop_region(asset, region)
                        pictures.append((cropped, self._as_text(region.get("label")), region.get("id")))
                    except (ValueError, TypeError, KeyError, OSError):
                        self.warnings.append(f"{self.source_labels[asset_id]} 的裁剪区域无效，已保留完整图片。")
                        pictures = []
                        break
                if not pictures:
                    pictures = [(asset["image_bytes"], "", None)]
                normalized = []
                for image_bytes, label, region_id in pictures:
                    with Image.open(io.BytesIO(image_bytes)) as image:
                        width, height = image.size
                    if len(batch) > 1:
                        label = self.source_labels[asset_id] + (f"　{label}" if label else "")
                    normalized.append({"bytes": image_bytes, "label": label, "region_id": region_id,
                                       "asset_id": asset_id, "width": width, "height": height,
                                       "aspect": width / height})
                source_pictures.append(normalized)
            # Interleave explicitly matched report sources; retain every source image.
            pictures = [source[i] for i in range(max(map(len, source_pictures)))
                        for source in source_pictures if i < len(source)]
            single_region_pages = any(self.analyses.get(link["asset_id"], {}).get("category")
                                      in ("genetic", "genetics", "genomic") for link in batch)
            page_groups = []
            while pictures:
                first = pictures.pop(0)
                page = [first]
                if pictures and not single_region_pages and first["aspect"] >= .65 and pictures[0]["aspect"] >= .65:
                    second = pictures[0]
                    # Do not mix a very wide text strip with a normal scan in narrow columns.
                    if (first["aspect"] >= 2.2) == (second["aspect"] >= 2.2):
                        page.append(pictures.pop(0))
                page_groups.append(page)
            category = categories[0] if len(set(categories)) == 1 else "检查资料"
            for page_index, page_pictures in enumerate(page_groups):
                slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
                title = category if len(page_groups) == 1 else f"{category}（{page_index + 1}/{len(page_groups)}）"
                self.add_header(slide, title)
                portrait = len(page_pictures) == 1 and page_pictures[0]["aspect"] < 1.15
                stacked = all(picture["aspect"] >= 2.2 for picture in page_pictures)
                for picture in page_pictures:
                    label_width = 7.0 if portrait else (11.7 if stacked or len(page_pictures) == 1 else 5.65)
                    label_limit = int(label_width * 72 / 8)
                    if len(self._wrap_lines(picture["label"], label_limit)) > 2:
                        units.append(picture["label"])
                        picture["label"] = ""
                if portrait:
                    picture = page_pictures[0]
                    self._textbox(slide, picture["label"], .8, 1.06, 7.0, .48, 14, self.C_ACC, True)
                    self._place_evidence_picture(slide, picture, .8, 1.55, 7.0, 5.25, 1)
                    text = self._take_complete_units(units, 34, 14)
                    self._textbox(slide, text, 8.1, 1.35, 4.4, 5.5, 18)
                elif stacked:
                    top = 1.08
                    for idx, picture in enumerate(page_pictures):
                        self._textbox(slide, picture["label"], .8, top, 11.7, .42, 14, self.C_ACC, True)
                        image_top = top + .5
                        height = self._place_evidence_picture(slide, picture, .8, image_top, 11.7,
                                                              1.9 if len(page_pictures) == 2 else 4.3, idx + 1)
                        top = image_top + height + .22
                    available = max(0, 6.8 - top)
                    text = self._take_complete_units(units, 80, int(available * 72 / 28))
                    if text:
                        self._textbox(slide, text, .8, top, 11.7, available, 18)
                else:
                    panel_width = 11.7 if len(page_pictures) == 1 else 5.65
                    for idx, picture in enumerate(page_pictures):
                        left = .8 + idx * 6.05
                        self._textbox(slide, picture["label"], left, 1.06, panel_width, .48, 14, self.C_ACC, True)
                        self._place_evidence_picture(slide, picture, left, 1.6, panel_width, 4.45, idx + 1)
                    text = self._take_complete_units(units, 90, 2)
                    if text:
                        self._textbox(slide, text, .8, 6.19, 11.7, .64, 16)
                page_asset_ids = [picture["asset_id"] for picture in page_pictures]
                self._evidence_page_footer(slide, page_asset_ids)
                for aid in dict.fromkeys(page_asset_ids):
                    self._record_evidence_slide(aid, links_by_id[aid],
                        [picture["region_id"] for picture in page_pictures if picture["asset_id"] == aid], "images")
            for text in self._detail_pages(units):
                slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
                self.add_header(slide, f"{category}（结果续）")
                self._textbox(slide, text, .8, 1.35, 11.7, 5.25, size=20)
                self._evidence_page_footer(slide, links_by_id)
                for aid, link in links_by_id.items():
                    self._record_evidence_slide(aid, link, [], "findings")

    def add_header(self, slide, text):
        if self.standard_case_theme:
            size = 32 if len(self._wrap_lines(text, 48)) == 1 else 24
            title = "\n".join(self._wrap_lines(text, 48 if size == 32 else 65))
            box = self._textbox(slide, title, .65, .12, 12.0, .72, size, RGBColor(0, 0, 0), True)
            box.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
            line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.65), Inches(.96), Inches(12.0), Inches(.025))
            line.fill.solid()
            line.fill.fore_color.rgb = self.C_PRI
            line.line.fill.background()
            return
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.33), Inches(0.9))
        shape.fill.solid()
        shape.fill.fore_color.rgb = self.C_PRI
        shape.line.fill.background()
        tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.05), Inches(12.3), Inches(0.8))
        tf = tb.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.word_wrap = True
        p = tb.text_frame.paragraphs[0]
        size = 28 if len(self._wrap_lines(text, 57)) == 1 else 18
        p.text = "\n".join(self._wrap_lines(text, 57 if size == 28 else 89))
        self._apply_font(p.font)
        p.font.size = Pt(size)
        p.line_spacing = 1.0
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)

    def make_cover(self):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.33), Inches(7.5))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(255, 255, 255) if self.standard_case_theme else self.C_PRI
        shape.line.fill.background()
        title = self._as_text(self.data.get("cover", {}).get("title", "病例汇报"))
        size = 48
        cover_detail = ""
        for candidate in (48, 40, 32, 28, 24, 20, 18):
            lines = self._wrap_lines(title, int(720 / (candidate * .55)))
            if len(lines) * candidate * 1.5 <= 4.8 * 72:
                size = candidate
                break
        else:
            cover_detail, title = title, "病例汇报"
        tb = slide.shapes.add_textbox(Inches(1.5), Inches(1.4), Inches(10), Inches(4.8))
        tb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tb.text_frame.paragraphs[0]
        p.text = "\n".join(self._wrap_lines(title, int(720 / (size * .55))))
        self._apply_font(p.font)
        p.font.size = Pt(size)
        p.font.bold = True
        p.font.color.rgb = RGBColor(0, 0, 0) if self.standard_case_theme else RGBColor(255, 255, 255)
        p.alignment = PP_ALIGN.CENTER
        if self.standard_case_theme:
            line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.5), Inches(5.9), Inches(10), Inches(.035))
            line.fill.solid()
            line.fill.fore_color.rgb = self.C_PRI
            line.line.fill.background()
        if cover_detail:
            self._make_body_pages("病例题目", [("", cover_detail, None, False)])

    def make_baseline(self):
        base_data = self.data.get("baseline", {})
        self._make_body_pages("病例介绍 (基线资料)", [
            ("【患者信息】", base_data.get("patient_info", ""), None, False),
            ("【主诉】", base_data.get("chief_complaint", ""), None, False),
            ("【临床诊断】", base_data.get("diagnosis", ""), None, False),
            ("【关键检查/病理】", base_data.get("key_exams", ""), None, False),
        ])

    def make_treatments(self):
        for treatment_index, tx in enumerate(self.data.get("treatments", [])):
            phase_name = self._phase_label(tx.get('phase', '阶段治疗'))
            self._make_body_pages(f"治疗经过：{phase_name}", [
                ("【治疗时间】", tx.get("duration", ""), self.C_PRI, True),
                ("【用药方案及局部治疗】", tx.get("regimen", ""), None, False),
                ("【影像学评估】", tx.get("imaging", ""), RGBColor(50, 50, 50), False),
                ("【肿瘤标志物】", tx.get("markers", ""), self.C_ACC, False),
            ])
            self.make_evidence("treatment", treatment_index)

    def make_current_admission(self):
        adm_data = self.data.get("current_admission")
        if not adm_data: return
        self._make_body_pages("本次入院评估及计划 (转归)", [
            ("【入院检验指标】", adm_data.get("exams", []), None, False),
            ("【影像学评估】", adm_data.get("imaging", ""), None, False),
            ("【后续计划】", adm_data.get("plan", []), None, False),
        ])

    @staticmethod
    def _timeline_style(event):
        kind = event.get("event_type", "Treatment")
        if kind in ("Plan", "Treatment Plan"):
            return RGBColor(110, 110, 110), True
        if kind == "Surgery":
            return RGBColor(70, 125, 56), False
        if kind == "Radiotherapy":
            return RGBColor(205, 113, 30), False
        if kind == "Treatment":
            return RGBColor(38, 91, 145), False
        return RGBColor(115, 21, 40), False

    def make_timeline_spec(self, spec):
        """Two editable vertical columns, following clinical case timeline references.

        Keep at most five nodes per column. Long source statements continue in
        further columns/pages, rather than shrinking the font or losing facts.
        """
        events = spec.get("events", self.data.get("timeline_events", []))
        if not isinstance(events, list) or not events:
            raise ValueError("时间轴需要至少一个有日期的事件。")
        if any(spec.get(key) for key in ("images", "sections", "bullets", "table")):
            raise ValueError("时间轴页不能同时包含图文或表格。")
        nodes = []
        for event in events:
            if not isinstance(event, dict) or not event.get("date") or not event.get("event"):
                raise ValueError("时间轴节点需要日期和事件。")
            color, planned = self._timeline_style(event)
            phase = self._phase_label(event.get("phase", ""))
            if planned and not re.search(r"计划|建议|拟|后续", phase):
                phase = "后续建议" + ("：" + phase if phase else "")
            phase_lines = self._wrap_lines(phase, 34) if phase else []
            event_lines = self._wrap_lines(self._as_text(event["event"]), 34)
            date_lines = self._wrap_lines(self._as_text(event["date"]), 15)
            # A long phase/date is also preserved; huge labels use normal body pages.
            if len(phase_lines) > 8 or len(date_lines) > 12:
                raise ValueError("时间轴日期或阶段名过长，请使用简短日期和阶段名称。")
            for offset in range(0, len(event_lines), 6):
                body = event_lines[offset:offset + 6]
                height = max(len(date_lines) * .26, len(phase_lines) * .25 + len(body) * .27) + .16
                nodes.append({"date": date_lines, "phase": phase_lines, "body": body,
                              "height": max(.92, height), "color": color, "planned": planned})
        subtitle = self._as_text(spec.get("subtitle"))
        subtitle_lines = self._wrap_lines(subtitle, 105) if subtitle else []
        if len(subtitle_lines) > 2:
            raise ValueError("时间轴副标题最多两行。")
        top = 1.28 + .25 * len(subtitle_lines)
        capacity = 6.79 - top
        columns, column, height = [], [], 0.0
        for node in nodes:
            if node["height"] > capacity:
                raise ValueError("时间轴节点无法在可读字号下排入。")
            if column and (len(column) == 5 or height + node["height"] > capacity):
                columns.append(column)
                column, height = [], 0.0
            column.append(node)
            height += node["height"]
        if column:
            columns.append(column)
        pages = math.ceil(len(columns) / 2)
        sources = self._plan_source_ids(spec)
        for page in range(pages):
            slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
            title = self._as_text(spec.get("title") or "诊疗时间轴")
            self.add_header(slide, title if pages == 1 else f"{title}（{page + 1}/{pages}）")
            if subtitle:
                self._textbox(slide, "\n".join(subtitle_lines), .65, 1.08, 12, .5, 14, RGBColor(100, 100, 100))
            for col_index, group in enumerate(columns[page * 2:page * 2 + 2]):
                left = .65 + col_index * 6.1
                axis_x = left + 1.64
                gap = min(.16, (capacity - sum(n["height"] for n in group)) / max(1, len(group) - 1))
                line = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(axis_x - .025),
                                             Inches(top + .13), Inches(.05), Inches(capacity - .17))
                line.fill.solid()
                line.fill.fore_color.rgb = RGBColor(185, 185, 185)
                line.line.fill.background()
                y = top
                for node in group:
                    color = node["color"]
                    self._textbox(slide, "\n".join(node["date"]), left, y + .02, 1.5,
                                  len(node["date"]) * .26, 16, color, True)
                    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(axis_x - .075),
                                                   Inches(y + .08), Inches(.15), Inches(.15))
                    circle.fill.solid()
                    circle.fill.fore_color.rgb = RGBColor(255, 255, 255) if node["planned"] else color
                    circle.line.color.rgb = color
                    circle.line.width = Pt(1.5)
                    if node["planned"]:
                        circle.line.dash_style = MSO_LINE_DASH_STYLE.DASH
                    body_y = y
                    for row in node["phase"]:
                        self._textbox(slide, row, left + 1.85, body_y, 3.85, .25, 15, color, True)
                        body_y += .25
                    for row in node["body"]:
                        self._textbox(slide, row, left + 1.85, body_y, 3.85, .27, 16)
                        body_y += .27
                    y += node["height"] + gap
            self._compact_source_footer(slide, sources, spec, [])

    def make_timeline(self):
        events = self.data.get("timeline_events", [])
        if not events: return
        # Six events per page keeps endpoint cards inside the slide and retains all events.
        for start in range(0, len(events), 6):
            self._make_timeline_page(events[start:start + 6], start // 6 + 1, math.ceil(len(events) / 6))

    def _make_timeline_page(self, events, page, total_pages):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        title = "全病程时间轴概览 (Timeline)"
        if total_pages > 1:
            title += f"（{page}/{total_pages}）"
        self.add_header(slide, title)
        line_y = Inches(4.2)
        start_x = Inches(1.55)
        total_width = 10.2
        count = len(events)
        card_width = Inches(2.6); card_height = Inches(1.65)
        font_size_date = Pt(12); font_size_body = Pt(11)

        main_line = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, start_x - Inches(0.2), line_y - Inches(0.05), Inches(total_width + 0.4), Inches(0.1))
        main_line.fill.solid()
        main_line.fill.fore_color.rgb = RGBColor(220, 220, 220)
        main_line.line.fill.background()

        for i, evt in enumerate(events):
            if count > 1: x = start_x + Inches(total_width * (i / (count - 1)))
            else: x = start_x + Inches(total_width / 2)
            event_text = evt.get("event", "")
            phase_text = self._phase_label(evt.get("phase", ""))
            event_type = evt.get("event_type", "Treatment")
            is_pd = "进展" in event_text or "PD" in event_text.upper() or "复发" in event_text
            is_control = "PR" in event_text.upper() or "SD" in event_text.upper() or "缩小" in event_text
            if is_pd: node_color = RGBColor(220, 50, 50)
            elif is_control and event_type == "Evaluation": node_color = RGBColor(46, 139, 87)
            else: node_color = self.C_PRI

            stem_height = Inches(1.0)
            stem_top = line_y - stem_height if i % 2 == 0 else line_y
            stem = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, stem_top, Inches(0.03), stem_height)
            stem.fill.solid()
            stem.fill.fore_color.rgb = node_color
            stem.line.fill.background()
            circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, x - Inches(0.15), line_y - Inches(0.15), Inches(0.3), Inches(0.3))
            circle.fill.solid()
            circle.fill.fore_color.rgb = node_color
            circle.line.color.rgb = RGBColor(255, 255, 255); circle.line.width = Pt(2)

            card_top = line_y - stem_height - card_height if i % 2 == 0 else line_y + stem_height
            card_x = x - (card_width / 2)
            card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, card_x, card_top, card_width, card_height)
            card.fill.solid()
            card.fill.fore_color.rgb = RGBColor(250, 250, 250)
            card.line.color.rgb = node_color; card.line.width = Pt(1.5)
            tf = card.text_frame
            tf.word_wrap = True
            tf.margin_left = Inches(0.05); tf.margin_right = Inches(0.05); tf.margin_top = Inches(0.05)
            p0 = tf.paragraphs[0]
            p0.text = evt.get("date", "")
            p0.font.bold = True; p0.font.size = font_size_date; p0.font.color.rgb = node_color; p0.alignment = PP_ALIGN.CENTER

            if phase_text and phase_text != "评估":
                p_phase = tf.add_paragraph()
                p_phase.text = f"【{phase_text}】"
                p_phase.font.size = font_size_body; p_phase.font.bold = True; p_phase.font.color.rgb = node_color; p_phase.alignment = PP_ALIGN.CENTER
            elif event_type == "Evaluation":
                p_phase = tf.add_paragraph()
                p_phase.text = "【疗效评估】"
                p_phase.font.size = font_size_body; p_phase.font.bold = True; p_phase.font.color.rgb = node_color; p_phase.alignment = PP_ALIGN.CENTER
            p1 = tf.add_paragraph()
            p1.text = event_text
            p1.font.size = font_size_body; p1.font.color.rgb = RGBColor(30, 30, 30); p1.alignment = PP_ALIGN.CENTER

    def make_summary(self):
        summary_data = self.data.get("summary", {})
        highlights = []
        discussion = []
        if isinstance(summary_data, list):
            highlights = summary_data
        elif isinstance(summary_data, dict):
            highlights = summary_data.get("highlights", [])
            discussion = summary_data.get("discussion", [])

        sections = [("", highlights, None, True)]
        if discussion:
            sections.append(("思考：", discussion, None, True))
        self._make_body_pages("病例思考与总结", sections)

    def build(self):
        plan = self.data.get("presentation_plan")
        if isinstance(plan, dict):
            planned_slides = plan.get("slides")
        elif isinstance(plan, list):
            planned_slides = plan
        else:
            planned_slides = None
        if planned_slides is not None:
            if not isinstance(planned_slides, list) or not planned_slides:
                raise ValueError("幻灯计划需要至少一页内容。")
        elif self.data.get("presentation_mode", "compact") != "detailed":
            planned_slides = self.compact_plan()
        if planned_slides is not None:
            for spec in planned_slides:
                self.make_planned_slide(spec)
            ppt_stream = io.BytesIO()
            self.prs.save(ppt_stream)
            ppt_stream.seek(0)
            return ppt_stream
        self.make_cover()
        self.make_baseline()
        self.make_evidence("baseline")
        self.make_treatments()
        self.make_current_admission()
        self.make_evidence("current_admission")
        self.make_timeline()
        self.make_evidence("appendix")
        self.make_summary()
        ppt_stream = io.BytesIO()
        self.prs.save(ppt_stream)
        ppt_stream.seek(0)
        return ppt_stream
